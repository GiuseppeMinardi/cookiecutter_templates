"""Test suite for calculate_hours module."""

from pathlib import Path
from io import StringIO
import sys

import pandas as pd
import pytest
from openpyxl import load_workbook


# We need to add the parent directory to sys.path to import the script
SCRIPT_DIR = (Path(__file__).parent.parent / "{{ cookiecutter.repo_name }}").resolve()
sys.path.insert(0, str(SCRIPT_DIR))

from calculate_hours import load_csv, process_data, modify_excel, main


@pytest.fixture
def sample_csv_file(tmp_path):
    """Create a sample CSV file for testing."""
    csv_content = """start,end,note
01/01/2024_09:00,01/01/2024_10:30,Morning meeting
01/01/2024_11:00,01/01/2024_13:00,Development work
01/01/2024_14:00,01/01/2024_16:30,Code review
"""
    csv_file = tmp_path / "test_log.csv"
    csv_file.write_text(csv_content)
    return csv_file


@pytest.fixture
def sample_dataframe():
    """Create a sample DataFrame for testing."""
    return pd.DataFrame({
        "start": ["01/01/2024_09:00", "01/01/2024_11:00"],
        "end": ["01/01/2024_10:30", "01/01/2024_13:00"],
        "note": ["Morning meeting", "Development"]
    })


class TestLoadCSV:
    """Test suite for load_csv function."""

    def test_load_csv_success(self, sample_csv_file):
        """Test successfully loading a CSV file."""
        df = load_csv(sample_csv_file)
        assert isinstance(df, pd.DataFrame)
        assert len(df) == 3
        assert list(df.columns) == ["start", "end", "note"]

    def test_load_csv_content(self, sample_csv_file):
        """Test that CSV content is correctly loaded."""
        df = load_csv(sample_csv_file)
        assert df.iloc[0]["start"] == "01/01/2024_09:00"
        assert df.iloc[0]["end"] == "01/01/2024_10:30"
        assert df.iloc[0]["note"] == "Morning meeting"

    def test_load_csv_empty_file(self, tmp_path):
        """Test loading a CSV file with only headers."""
        csv_file = tmp_path / "empty.csv"
        csv_file.write_text("start,end,note\n")
        df = load_csv(csv_file)
        assert len(df) == 0
        assert list(df.columns) == ["start", "end", "note"]

    def test_load_csv_file_not_found(self, tmp_path):
        """Test error handling for non-existent file."""
        non_existent = tmp_path / "non_existent.csv"
        with pytest.raises(FileNotFoundError):
            load_csv(non_existent)


class TestProcessData:
    """Test suite for process_data function."""

    def test_process_data_creates_required_columns(self, sample_dataframe):
        """Test that processing creates duration_hrs and cost_eur columns."""
        df = process_data(sample_dataframe)
        assert "duration_hrs" in df.columns
        assert "cost_eur" in df.columns

    def test_process_data_duration_calculation(self, sample_dataframe):
        """Test correct duration calculation in hours."""
        df = process_data(sample_dataframe)
        # First entry: 09:00 to 10:30 = 1.5 hours
        assert abs(df.iloc[0]["duration_hrs"] - 1.5) < 0.01
        # Second entry: 11:00 to 13:00 = 2 hours
        assert abs(df.iloc[1]["duration_hrs"] - 2.0) < 0.01

    def test_process_data_cost_calculation(self, sample_dataframe):
        """Test correct cost calculation with custom hourly fee."""
        df = process_data(sample_dataframe, hourly_fee=30.0)
        # First entry: 1.5 hours * 30 EUR = 45 EUR
        assert abs(df.iloc[0]["cost_eur"] - 45.0) < 0.01
        # Second entry: 2 hours * 30 EUR = 60 EUR
        assert abs(df.iloc[1]["cost_eur"] - 60.0) < 0.01

    def test_process_data_default_hourly_fee(self, sample_dataframe):
        """Test that default hourly fee is 25.0 EUR."""
        df = process_data(sample_dataframe)
        # First entry: 1.5 hours * 25 EUR = 37.5 EUR
        assert abs(df.iloc[0]["cost_eur"] - 37.5) < 0.01

    def test_process_data_custom_date_format(self, tmp_path):
        """Test processing with custom date format."""
        df = pd.DataFrame({
            "start": ["2024-01-01 09:00", "2024-01-01 11:00"],
            "end": ["2024-01-01 10:30", "2024-01-01 13:00"],
            "note": ["Meeting", "Work"]
        })
        processed = process_data(df, date_format="%Y-%m-%d %H:%M")
        assert abs(processed.iloc[0]["duration_hrs"] - 1.5) < 0.01
        assert abs(processed.iloc[1]["duration_hrs"] - 2.0) < 0.01

    def test_process_data_converts_to_datetime(self, sample_dataframe):
        """Test that start and end are converted to datetime objects."""
        df = process_data(sample_dataframe)
        assert pd.api.types.is_datetime64_any_dtype(df["start"])
        assert pd.api.types.is_datetime64_any_dtype(df["end"])

    def test_process_data_preserves_note_column(self, sample_dataframe):
        """Test that note column is preserved during processing."""
        df = process_data(sample_dataframe)
        assert "note" in df.columns
        assert df.iloc[0]["note"] == "Morning meeting"


class TestModifyExcel:
    """Test suite for modify_excel function."""

    def test_modify_excel_file_exists(self, sample_csv_file, tmp_path):
        """Test that modify_excel doesn't fail on valid file."""
        output_file = tmp_path / "output.xlsx"
        df = load_csv(sample_csv_file)
        processed = process_data(df)
        processed = processed[
            ["start", "end", "duration_hrs", "cost_eur", "note"]
        ].rename(
            columns={
                "start": "Start Time",
                "end": "End Time",
                "duration_hrs": "Duration (hrs)",
                "cost_eur": "Cost (€)",
                "note": "Note",
            }
        )
        processed.to_excel(output_file, index=False)
        
        # This should not raise an error
        modify_excel(output_file)
        assert output_file.exists()

    def test_modify_excel_preserves_data(self, sample_csv_file, tmp_path):
        """Test that modify_excel preserves the data in the file."""
        output_file = tmp_path / "output.xlsx"
        df = load_csv(sample_csv_file)
        processed = process_data(df)
        processed = processed[
            ["start", "end", "duration_hrs", "cost_eur", "note"]
        ].rename(
            columns={
                "start": "Start Time",
                "end": "End Time",
                "duration_hrs": "Duration (hrs)",
                "cost_eur": "Cost (€)",
                "note": "Note",
            }
        )
        processed.to_excel(output_file, index=False)
        modify_excel(output_file)
        
        # Read back and verify data
        wb = load_workbook(output_file)
        ws = wb.active
        assert ws["A1"].value == "Start Time"
        assert ws["B1"].value == "End Time"
        assert ws["C1"].value == "Duration (hrs)"
        assert ws["D1"].value == "Cost (€)"
        assert ws["E1"].value == "Note"

    def test_modify_excel_has_summary_section(self, sample_csv_file, tmp_path):
        """Test that modify_excel creates summary section."""
        output_file = tmp_path / "output.xlsx"
        df = load_csv(sample_csv_file)
        processed = process_data(df)
        processed = processed[
            ["start", "end", "duration_hrs", "cost_eur", "note"]
        ].rename(
            columns={
                "start": "Start Time",
                "end": "End Time",
                "duration_hrs": "Duration (hrs)",
                "cost_eur": "Cost (€)",
                "note": "Note",
            }
        )
        processed.to_excel(output_file, index=False)
        modify_excel(output_file)
        
        wb = load_workbook(output_file)
        ws = wb.active
        # Check for summary labels
        assert ws["G2"].value == "TOTAL DUE:"

    def test_modify_excel_has_total_formulas(self, sample_csv_file, tmp_path):
        """Test that modify_excel adds total formulas."""
        output_file = tmp_path / "output.xlsx"
        df = load_csv(sample_csv_file)
        processed = process_data(df)
        processed = processed[
            ["start", "end", "duration_hrs", "cost_eur", "note"]
        ].rename(
            columns={
                "start": "Start Time",
                "end": "End Time",
                "duration_hrs": "Duration (hrs)",
                "cost_eur": "Cost (€)",
                "note": "Note",
            }
        )
        processed.to_excel(output_file, index=False)
        modify_excel(output_file)
        
        wb = load_workbook(output_file)
        ws = wb.active
        # Check for SUM formulas
        assert ws["H2"].value.startswith("=SUM(C")
        assert ws["I2"].value.startswith("=SUM(D")


class TestMain:
    """Test suite for main function."""

    def test_main_creates_output_file(self, sample_csv_file, tmp_path):
        """Test that main function creates output file."""
        output_file = tmp_path / "output.xlsx"
        main(
            input_csv_path=sample_csv_file,
            output_excel_path=output_file,
            date_format="%d/%m/%Y_%H:%M",
            hourly_fee=25.0,
        )
        assert output_file.exists()

    def test_main_output_has_data(self, sample_csv_file, tmp_path):
        """Test that main function outputs correct data."""
        output_file = tmp_path / "output.xlsx"
        main(
            input_csv_path=sample_csv_file,
            output_excel_path=output_file,
            date_format="%d/%m/%Y_%H:%M",
            hourly_fee=25.0,
        )
        
        df = pd.read_excel(output_file)
        assert len(df) == 3
        assert "Duration (hrs)" in df.columns
        assert "Cost (€)" in df.columns

    def test_main_custom_hourly_fee(self, sample_csv_file, tmp_path):
        """Test main function with custom hourly fee."""
        output_file = tmp_path / "output.xlsx"
        main(
            input_csv_path=sample_csv_file,
            output_excel_path=output_file,
            date_format="%d/%m/%Y_%H:%M",
            hourly_fee=50.0,
        )
        
        df = pd.read_excel(output_file)
        # First row: 1.5 hours * 50 EUR = 75 EUR
        assert abs(df.iloc[0]["Cost (€)"] - 75.0) < 0.01

    def test_main_with_single_entry(self, tmp_path):
        """Test main function with single entry."""
        csv_file = tmp_path / "single.csv"
        csv_file.write_text("start,end,note\n01/01/2024_09:00,01/01/2024_10:00,Work\n")
        
        output_file = tmp_path / "output.xlsx"
        main(
            input_csv_path=csv_file,
            output_excel_path=output_file,
            date_format="%d/%m/%Y_%H:%M",
            hourly_fee=20.0,
        )
        
        df = pd.read_excel(output_file)
        assert len(df) == 1
        assert abs(df.iloc[0]["Duration (hrs)"] - 1.0) < 0.01
        assert abs(df.iloc[0]["Cost (€)"] - 20.0) < 0.01


class TestEdgeCases:
    """Test suite for edge cases and error conditions."""

    def test_process_data_with_zero_duration(self):
        """Test processing data with identical start and end times."""
        df = pd.DataFrame({
            "start": ["01/01/2024_09:00"],
            "end": ["01/01/2024_09:00"],
            "note": ["Quick task"]
        })
        processed = process_data(df)
        assert abs(processed.iloc[0]["duration_hrs"] - 0.0) < 0.01
        assert abs(processed.iloc[0]["cost_eur"] - 0.0) < 0.01

    def test_process_data_with_zero_hourly_fee(self, sample_dataframe):
        """Test processing data with zero hourly fee."""
        df = process_data(sample_dataframe, hourly_fee=0.0)
        assert all(df["cost_eur"] == 0.0)

    def test_process_data_preserves_original_columns(self, sample_dataframe):
        """Test that original columns are preserved."""
        df = process_data(sample_dataframe)
        assert "start" in df.columns
        assert "end" in df.columns
        assert "note" in df.columns
